import csv
import io
from datetime import datetime, timedelta

import pytest
from conftest import ADMIN_HEADERS, START
from openpyxl import load_workbook


def end_after(seconds: float) -> str:
    return (datetime.fromisoformat(START) + timedelta(seconds=seconds)).isoformat()


def read_csv(response) -> list[list[str]]:
    return list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))


def read_xlsx(response):
    return load_workbook(io.BytesIO(response.content)).active


@pytest.fixture
async def scenario(make_team, make_driver, make_challenge, make_penalty_type, make_attempt, competition_client):
    challenge = await make_challenge()
    penalty_type = await make_penalty_type(amount=10)
    a = await make_team("=1+1")  # must stay text in exports
    b = await make_team("B", category="professional_class")
    driver_a = await make_driver(a["id"], name="Ä driver")
    driver_b = await make_driver(b["id"])
    fast = await make_attempt(a["id"], driver_a["id"], challenge, end_time=end_after(60),
                              penalty_type=penalty_type, penalty_count=1)
    invalid = await make_attempt(a["id"], driver_a["id"], challenge, end_time=end_after(50))
    await competition_client.patch(
        f"/attempts/{invalid['id']}/validity", headers=ADMIN_HEADERS, json={"is_valid": False}
    )
    await make_attempt(b["id"], driver_b["id"], challenge, end_time=end_after(90))
    return {"challenge": challenge, "fast": fast["id"], "invalid": invalid["id"]}


@pytest.mark.asyncio
async def test_leaderboard_csv_has_the_leaderboard_rows(competition_client, scenario):
    response = await competition_client.get(f"/export/leaderboard/{scenario['challenge']}/category/close_to_series")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    disposition = f'attachment; filename="leaderboard_challenge{scenario["challenge"]}_close_to_series.csv"'
    assert response.headers["content-disposition"] == disposition
    header, *rows = read_csv(response)
    assert header == ["rank", "team", "category", "score", "attempt_id", "time_s", "energy_used"]
    assert [row[:5] for row in rows] == [["1", "=1+1", "close_to_series", "100.0", str(scenario["fast"])]]


@pytest.mark.asyncio
async def test_leaderboard_xlsx_keeps_text_as_text(competition_client, scenario):
    response = await competition_client.get(
        f"/export/leaderboard/{scenario['challenge']}/category/close_to_series?format=xlsx"
    )

    sheet = read_xlsx(response)
    assert response.headers["content-disposition"].endswith('.xlsx"')
    assert sheet["B2"].value == "=1+1"
    assert sheet["B2"].data_type == "s"
    assert sheet["D2"].value == 100.0


@pytest.mark.asyncio
async def test_attempts_export_has_all_attempts_with_penalties(competition_client, scenario):
    response = await competition_client.get(f"/export/attempts/{scenario['challenge']}")

    header, *rows = read_csv(response)
    by_id = {int(row[0]): dict(zip(header, row)) for row in rows}
    assert len(rows) == 3
    assert by_id[scenario["fast"]]["driver"] == "Ä driver"
    assert float(by_id[scenario["fast"]]["penalty_s"]) == 10
    assert float(by_id[scenario["fast"]]["time_with_penalties_s"]) == pytest.approx(70, abs=1e-3)
    assert by_id[scenario["invalid"]]["is_valid"] == "False"


@pytest.mark.asyncio
async def test_attempts_export_category_filter_and_xlsx(competition_client, scenario):
    response = await competition_client.get(
        f"/export/attempts/{scenario['challenge']}?category=professional_class&format=xlsx"
    )

    sheet = read_xlsx(response)
    assert response.headers["content-disposition"].endswith(
        f'attempts_challenge{scenario["challenge"]}_professional_class.xlsx"'
    )
    assert [row[1] for row in sheet.iter_rows(min_row=2, values_only=True)] == ["B"]


@pytest.mark.asyncio
async def test_export_unknown_challenge_is_404(competition_client):
    assert (await competition_client.get("/export/attempts/999")).status_code == 404
    assert (await competition_client.get("/export/leaderboard/999/category/close_to_series")).status_code == 404
